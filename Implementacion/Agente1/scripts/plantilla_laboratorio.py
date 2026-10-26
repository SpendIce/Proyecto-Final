"""Utilidades para volcar los anexos del laboratorio sobre las plantillas oficiales.

Las plantillas `3.1 - Anexo 1 - Presentación 2025.docx`, `3.2 - Anexo 2 -
Calidad 2025.docx` y `3.3 - Anexo 3 - Seguridad 2025.docx` viven en
`Contenido/DocumentacionAnteproyecto/Revision-Laboratorio-2025/` y son el
formato que pide el laboratorio (issue #57): se entregan en Word **sobre esas
plantillas**, no sobre una tabla propia.

El reparto de columnas lo fija el correo del laboratorio del 14/10 y las notas
al pie de cada plantilla («El Equipo de Proyecto completará esta tabla
consignando los factores y atributos»):

- **El equipo de proyecto** completa los ítems: `Factor de calidad` + `Atributo`
  en el Anexo 2, `Componente` + `Ítem` en el Anexo 3.
- **El laboratorio** completa las últimas columnas (nivel de competencia y
  comentarios), y por eso pide Word.

Las plantillas son las del ciclo 2025: el encabezado de página dice «Ciclo
Lectivo 2025» y se actualiza al ciclo vigente (`CICLO_LECTIVO`).

Este módulo sólo aporta la plomería de `python-docx`: abrir la plantilla,
reemplazar marcadores y armar las filas de la tabla. Lo usan los tres scripts
`anexo*_*.py`.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

# Título con el que se identifica el proyecto en la cabecera de cada plantilla.
# Es el mismo texto que usa el Anexo 1, para no tener dos títulos del proyecto.
TITULO_PROYECTO = (
    "Proyecto Centenario, Agente 1: Extensión Bot para la Secretaría de "
    "Extensión Universitaria de la FIE (Facultad de Ingeniería del Ejército, "
    "Universidad de la Defensa Nacional)"
)


# Ciclo lectivo de la entrega. Las plantillas oficiales son las de 2025.
CICLO_LECTIVO = "2026"


def abrir_plantilla(ruta: Path) -> Any:
    """Abre una plantilla oficial en .docx y devuelve el documento."""

    try:
        from docx import Document
    except ImportError:
        raise SystemExit(
            "python-docx no está instalado; correr con "
            "`uv run --with python-docx python scripts/<script> --plantilla … --docx …`"
        )
    if not ruta.exists():
        raise SystemExit(f"no está la plantilla oficial: {ruta}")
    return Document(str(ruta))


def _fijar_texto(parrafo: Any, texto: str) -> None:
    """Deja el párrafo con `texto`, conservando el formato del primer run."""

    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        parrafo.add_run(texto)


def actualizar_ciclo_lectivo(doc: Any) -> None:
    """Cambia «Ciclo Lectivo 2025» del encabezado de página por el ciclo vigente.

    En la plantilla el año está partido en dos runs («Ciclo Lectivo 202» y
    «5»), así que se reescribe desde el run donde empieza la frase, con el
    formato de ese run. Sólo se mira el encabezado principal de cada sección:
    leer los de primera página o de páginas pares los crearía vacíos.
    """

    cambiados = 0
    for seccion in doc.sections:
        for parrafo in seccion.header.paragraphs:
            if "Ciclo Lectivo" not in parrafo.text:
                continue
            for indice, run in enumerate(parrafo.runs):
                if "Ciclo" in run.text:
                    resto = "".join(r.text for r in parrafo.runs[indice:])
                    run.text = re.sub(
                        r"Ciclo Lectivo \d{4}", f"Ciclo Lectivo {CICLO_LECTIVO}", resto
                    )
                    for siguiente in parrafo.runs[indice + 1 :]:
                        siguiente.text = ""
                    cambiados += 1
                    break
    if not cambiados:
        raise SystemExit("la plantilla no tiene «Ciclo Lectivo» en el encabezado")


def buscar_parrafo(doc: Any, texto: str) -> Any:
    """Devuelve el primer párrafo que contiene `texto`."""

    for parrafo in doc.paragraphs:
        if texto in parrafo.text:
            return parrafo
    raise SystemExit(f"la plantilla no tiene el párrafo esperado: {texto!r}")


def reemplazar_parrafo(doc: Any, antiguo: str, nuevo: str) -> None:
    """Reemplaza el primer párrafo que contiene `antiguo`."""

    for parrafo in doc.paragraphs:
        if antiguo in parrafo.text:
            _fijar_texto(parrafo, nuevo)
            return
    raise SystemExit(f"la plantilla no tiene el párrafo esperado: {antiguo!r}")


def reemplazar_en_celdas(doc: Any, antiguo: str, nuevo: str) -> None:
    """Reemplaza `antiguo` en el primer celda de tabla que lo contiene.

    Sirve para el título de la tabla (`Factores de Calidad <Citar estándar
    aplicado>`), que es una celda combinada.
    """

    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if antiguo in celda.text:
                    _fijar_texto(celda.paragraphs[0], nuevo)
                    return
    raise SystemExit(f"la plantilla no tiene la celda esperada: {antiguo!r}")


def conservar_filas(tabla: Any, cantidad: int) -> None:
    """Borra de la tabla todas las filas salvo las primeras `cantidad`."""

    for fila in list(tabla.rows)[cantidad:]:
        tabla._tbl.remove(fila._tr)


def agregar_fila(tabla: Any, valores: list[str]) -> None:
    """Agrega una fila con `valores` y vacía las celdas que no se declaran."""

    celdas = tabla.add_row().cells
    if len(valores) > len(celdas):
        raise ValueError(
            f"la fila tiene {len(valores)} valores para {len(celdas)} columnas"
        )
    for i, celda in enumerate(celdas):
        celda.text = valores[i] if i < len(valores) else ""


def guardar(doc: Any, destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))


def verificar_columnas(tabla: Any, esperadas: list[str]) -> None:
    """Falla si la fila de encabezado de la plantilla no es la esperada.

    Protege contra una plantilla oficial distinta a la del 2025: si cambian las
    columnas, el volcado tiene que fallar ruidosamente y no escribir corrido.
    La comparación es por contención porque hay celdas de encabezado partidas en
    dos líneas («Atributo» + «(si es que hubiere para el proyecto)»).
    """

    for indice, fila in enumerate(tabla.rows):
        textos = [c.text.strip() for c in fila.cells]
        if len(textos) >= len(esperadas) and all(
            esperado in textos[i] for i, esperado in enumerate(esperadas)
        ):
            return
        if indice >= 3:
            break
    raise SystemExit(
        "la plantilla no tiene la fila de encabezado esperada: " + " | ".join(esperadas)
    )
